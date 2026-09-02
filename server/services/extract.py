"""文档/网页文本提取：PDF/DOCX/TXT/图片 OCR（预留）"""
from pathlib import Path

def extract_pdf_text(path: str) -> str:
    try:
        import PyPDF2
        reader = PyPDF2.PdfReader(path)
        parts = []
        for page in reader.pages:
            try:
                text = page.extract_text()
                if text:
                    parts.append(text)
            except Exception:
                pass
        return "\n".join(parts).strip()
    except Exception as e:
        return f"[PDF 提取失败：{e}]"


def extract_docx_text(path: str) -> str:
    try:
        import docx
        doc = docx.Document(path)
        parts = []
        for para in doc.paragraphs:
            if para.text:
                parts.append(para.text)
        # 简单处理表格文本
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text for cell in row.cells if cell.text]
                if cells:
                    parts.append("\t".join(cells))
        return "\n".join(parts).strip()
    except Exception as e:
        return f"[DOCX 提取失败：{e}]"


def extract_text(path: str) -> str:
    p = Path(path)
    if not p.exists():
        return "[文件不存在]"
    ext = p.suffix.lower()
    if ext == '.pdf':
        return extract_pdf_text(path)
    if ext in ('.docx', '.doc'):
        # python-docx 不支持 .doc；.doc 尝试按 docx 读，失败再提示
        try:
            return extract_docx_text(path)
        except Exception as e:
            return f"[DOC 提取失败：{e}]"
    if ext in ('.txt', '.md', '.html', '.htm', '.csv'):
        try:
            return p.read_text(encoding='utf-8', errors='ignore')
        except Exception as e:
            return f"[文本读取失败：{e}]"
    if ext in ('.png', '.jpg', '.jpeg', '.gif', '.webp'):
        return "[图片内容暂不支持自动 OCR，请在上传处填写文字说明或换用文本/链接类型]"
    return f"[暂不支持的格式 {ext}]"
